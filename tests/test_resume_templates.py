"""ATS-friendly resume templates: recommendation, order, rendering and the API."""
from pathlib import Path
from uuid import uuid4

import pytest

from job_hunting_agent import web
from job_hunting_agent.ats import score_resume
from job_hunting_agent.db import record_run, run_for_user
from job_hunting_agent.models import CandidateProfile
from job_hunting_agent.resume import parse_resume
from job_hunting_agent.resume_builder import write_base_resume
from job_hunting_agent.resume_templates import TEMPLATES, apply_template_order, catalog, get_template, recommend_template
from tests.test_web import authenticated_client

RESUME_TEXT = """Agnirudra Banerjee
Kolkata, India | +91 98765 43210 | agni@example.com | linkedin.com/in/agni
SUMMARY
Data engineer with 3 years of experience building batch and streaming pipelines on AWS with Python, Spark and Airflow.
SKILLS
Languages: Python, SQL, Scala
Data: Apache Spark, Kafka, Airflow, dbt, Snowflake
EXPERIENCE
Data Engineer | Kairoz Technologies | Kolkata | Jan 2023 - Present
- Built Spark pipelines processing 40 million IoT events per day, cutting batch time by 35%.
- Migrated 12 Airflow DAGs to managed Airflow, reducing failed runs from 9% to 2%.
PROJECTS
Resume Ranking Engine - Python and scikit-learn service that ranks resumes against job descriptions.
EDUCATION
B.Tech in Computer Science (AI-ML) - Sister Nivedita University, 2021 - CGPA 8.4
CERTIFICATIONS
AWS Certified Data Engineer - Associate (2024)
"""
PROFILE = CandidateProfile(name="Agnirudra Banerjee", email="agni@example.com", phone="+91 98765 43210",
                           target_roles=("Data Engineer",), skills=("Python", "SQL", "Spark"), experience_years=3)


@pytest.mark.parametrize("years, role, expected", [
    (0, "Data Engineering", "graduate"), (1.5, "Sales and Marketing", "graduate"),
    (3, "Data Engineering", "technical"), (4, "AI and Machine Learning", "technical"),
    (4, "Sales and Marketing", "modern"), (5, "Finance and Investment", "professional"),
    (9, "Data Engineering", "executive"),
])
def test_agent_recommends_a_template_from_experience_and_role(years, role, expected):
    key, reason = recommend_template(years, role)
    assert key == expected and reason
    listed = catalog(key, reason)
    assert listed[0]["key"] == expected and listed[0]["recommended"] and listed[0]["reason"] == reason
    assert sum(item["recommended"] for item in listed) == 1 and len(listed) == len(TEMPLATES)


def test_templates_reorder_sections_without_dropping_any():
    sections = [("CONTACT", ["c"]), ("PROFESSIONAL SUMMARY", ["s"]), ("TECHNICAL SKILLS", ["k"]), ("PROFESSIONAL EXPERIENCE", ["e"]),
                ("PROJECTS", ["p"]), ("EDUCATION", ["d"]), ("CERTIFICATIONS", ["x"]), ("ACHIEVEMENTS", ["a"]), ("LANGUAGES", ["l"])]
    for template in TEMPLATES.values():
        ordered = apply_template_order(sections, template)
        assert sorted(ordered) == sorted(sections), template.key
        assert [heading for heading, _ in ordered][:2] == ["CONTACT", "PROFESSIONAL SUMMARY"]
    order = lambda key: [heading for heading, _ in apply_template_order(sections, get_template(key))]
    assert order("graduate").index("EDUCATION") < order("graduate").index("PROFESSIONAL EXPERIENCE")
    assert order("technical").index("TECHNICAL SKILLS") < order("technical").index("PROFESSIONAL EXPERIENCE")
    assert order("executive").index("PROFESSIONAL EXPERIENCE") < order("executive").index("TECHNICAL SKILLS")
    assert order("professional") == [heading for heading, _ in sections]
    assert get_template("modern").heading("PROFESSIONAL EXPERIENCE") == "Professional Experience"
    assert get_template("unknown").key == "professional"


@pytest.mark.parametrize("key", list(TEMPLATES))
def test_every_template_builds_an_ats_safe_validated_resume(key, tmp_path, monkeypatch):
    from docx import Document

    # The PDF fallback renderer keeps the test fast; LibreOffice rendering is covered elsewhere.
    monkeypatch.setattr("job_hunting_agent.resume_builder.convert_docx_to_pdf", lambda *_: False)
    source = tmp_path / "resume.txt"
    source.write_text(RESUME_TEXT)
    resume = parse_resume(str(source))
    artifact = write_base_resume(resume, score_resume(resume, PROFILE), PROFILE, tmp_path / "out", template=key)
    template = TEMPLATES[key]
    validation = artifact["validation"]
    assert artifact["template"] == key and artifact["template_name"] == template.name
    assert validation["passed"], validation["issues"]
    assert validation["page_count"] == 1 and validation["pdf_text_coverage"] >= 0.95
    document = Document(artifact["docx_path"])
    assert not document.tables and not document.inline_shapes  # no tables, images or icons
    assert len(document.sections) == 1
    assert document.styles["Normal"].font.name == template.font
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for fact in ("Agnirudra Banerjee", "40 million IoT events", "Sister Nivedita University", "AWS Certified Data Engineer"):
        assert fact in text
    assert template.heading("PROFESSIONAL EXPERIENCE") in text
    name = document.paragraphs[0]
    assert (name.style.paragraph_format.alignment == 1) == (template.header_align == "center")
    if key != "professional":
        assert template.name.replace(" / ", "-").replace(" ", "-") in Path(artifact["docx_path"]).name


def _run_with_template_artifact(email: str, tmp_path: Path) -> int:
    docx, pdf, png = tmp_path / "r.docx", tmp_path / "r.pdf", tmp_path / "p-1.png"
    for path in (docx, pdf, png):
        path.write_bytes(b"x")
    return record_run(email, {
        "trigger": "manual", "generated_at": "2026-10-31T10:00:00", "ats_report": {"score": 80}, "jobs": [],
        "application_summary": {}, "output_dir": "",
        "improved_resume": {
            "status": "ready", "template": "technical", "recommended_template": "technical",
            "template_reason": "For Data Engineering roles, recruiters scan your skills and projects first.",
            "docx_path": str(docx), "pdf_path": str(pdf), "tailored_resumes": [],
            "validation": {"passed": True, "page_count": 1, "rendered_pages": [str(png)]},
        },
    })


def test_template_api_lists_prepares_downloads_and_sets_the_tailored_template(monkeypatch, tmp_path):
    email = f"tpl-{uuid4().hex}@example.com"
    run_id = _run_with_template_artifact(email, tmp_path)
    submitted = []
    monkeypatch.setattr(web, "_submit_document_task", lambda key, *args: submitted.append(key) or True)
    client = authenticated_client(email)

    listing = client.get(f"/api/resume-templates?run_id={run_id}").json()
    assert listing["recommended"] == "technical" and listing["templates"][0]["key"] == "technical"
    assert listing["statuses"]["technical"] == {"status": "ready", "passed": True, "page_count": 1, "base": True}

    # The main resume already is the technical template: nothing to build.
    assert client.post(f"/api/runs/{run_id}/resume-templates/technical/prepare").json()["status"] == "ready"
    assert client.get(f"/api/runs/{run_id}/resume-templates/technical/docx").status_code == 200
    assert client.get(f"/api/runs/{run_id}/resume-templates/technical/preview").headers["content-type"] == "image/png"
    # Another template is built in the background, once.
    assert client.post(f"/api/runs/{run_id}/resume-templates/modern/prepare").json()["status"] == "preparing"
    assert submitted == [f"template:{run_id}:modern"]
    assert client.get(f"/api/runs/{run_id}/resume-status").json()["templates"]["modern"]["status"] == "preparing"
    assert client.get(f"/api/runs/{run_id}/resume-templates/modern/pdf").status_code == 404
    assert client.post(f"/api/runs/{run_id}/resume-templates/fancy/prepare").status_code == 404

    chosen = client.post(f"/api/runs/{run_id}/resume-template", json={"template": "executive"}).json()
    assert chosen == {"preferred_template": "executive", "name": "Executive"}
    status = client.get(f"/api/runs/{run_id}/resume-status").json()
    assert status["template"] == "technical" and status["preferred_template"] == "executive"
    assert client.post(f"/api/runs/{run_id}/resume-template", json={"template": "fancy"}).status_code == 400
    # Other users cannot see or change this run.
    other = authenticated_client(f"tpl-other-{uuid4().hex}@example.com")
    assert other.post(f"/api/runs/{run_id}/resume-templates/modern/prepare").status_code == 404
    assert other.get(f"/api/runs/{run_id}/resume-templates/technical/docx").status_code == 404


def test_template_worker_stores_the_artifact_and_reports_failures(monkeypatch, tmp_path):
    email = f"tpl-worker-{uuid4().hex}@example.com"
    run_id = _run_with_template_artifact(email, tmp_path)
    from dataclasses import asdict

    config = web._config_from_snapshot({"profile": asdict(PROFILE)})
    monkeypatch.setattr(web, "_document_inputs", lambda payload: (None, None, config, "hash"))
    monkeypatch.setattr(web, "cached_document", lambda data_dir, key, builder: ({
        "docx_path": "m.docx", "pdf_path": "m.pdf", "template": "modern",
        "validation": {"passed": True, "page_count": 1}}, False))
    monkeypatch.setattr(web, "_publish_document_artifact", lambda artifact, *args: artifact)
    web._prepare_template_document(email, run_id, "modern")
    improved = run_for_user(email, run_id)["payload"]["improved_resume"]
    assert improved["template_resumes"]["modern"]["template"] == "modern"
    assert web._template_statuses(improved)["modern"]["status"] == "ready"

    def broken(*args):
        raise RuntimeError("render failed")
    monkeypatch.setattr(web, "cached_document", broken)
    web._prepare_template_document(email, run_id, "graduate")
    improved = run_for_user(email, run_id)["payload"]["improved_resume"]
    assert improved["template_statuses"]["graduate"] == {"status": "failed", "error": "render failed"}


def test_client_payload_never_exposes_template_file_paths():
    payload = web._client_run_payload({"jobs": [], "applications": [], "improved_resume": {
        "status": "ready", "template": "modern", "template_resumes": {"graduate": {"docx_path": "/secret/a.docx", "validation": {}}},
    }}, 1)
    assert "template_resumes" not in payload["improved_resume"]
    assert payload["improved_resume"]["template_statuses"]["graduate"]["status"] == "ready"
