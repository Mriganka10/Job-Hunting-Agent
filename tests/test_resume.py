from pathlib import Path
import sys
from types import SimpleNamespace

from job_hunting_agent.resume import _read_pdf, detect_sections, extract_roles, extract_sections, extract_skills, normalize_text, parse_resume


def test_parse_text_resume_extracts_skills(tmp_path: Path) -> None:
    resume_path = tmp_path / "resume.txt"
    resume_path.write_text("Python Developer with SQL, AWS, and FastAPI experience.", encoding="utf-8")

    resume = parse_resume(resume_path)

    assert "Python" in resume.inferred_skills
    assert "SQL" in resume.inferred_skills
    assert "Python Developer" in resume.inferred_roles


def test_detect_sections_handles_common_resume_heading_variants() -> None:
    text = """PROFILE SUMMARY
Engineer
CORE SKILLS: Python and SQL
INTERNSHIPS
Worked on APIs
ACADEMIC PROJECTS
Built a reporting tool
ACADEMIC QUALIFICATIONS
B.Tech
"""

    assert set(detect_sections(text)) == {"summary", "skills", "experience", "projects", "education"}


def test_normalize_text_repairs_pdf_line_hyphenation_and_unicode() -> None:
    text = "Machine Learn-\ning\u00a0Engineer\n\n\nSkills"

    assert normalize_text(text) == "Machine Learning Engineer\nSkills"


def test_normalize_text_repairs_character_spaced_pdf_glyphs() -> None:
    text = "M A C H I N E  L E A R N I N G\nP y t h o n ,  S Q L"

    assert normalize_text(text) == "Machine Learning\nPython, SQL"


def test_parse_resume_canonicalizes_common_ats_keyword_formatting(tmp_path: Path) -> None:
    resume_path = tmp_path / "resume.txt"
    resume_path.write_text(
        "Skills: Aws, Gitlab, Fastapi, SQL, highquality data checks.",
        encoding="utf-8",
    )

    resume = parse_resume(resume_path)

    assert "AWS" in resume.inferred_skills
    assert "GitLab" in resume.inferred_skills
    assert "FastAPI" in resume.inferred_skills
    assert "high-quality data checks" in resume.text


def test_extract_sections_recovers_education_before_two_column_heading() -> None:
    text = """CANDIDATE NAME
candidate@example.com
2022 - 2026
CGPA - 9.38
B.Tech in CSE
Example University
EDUCATION
PROJECTS
Built an analytics project.
SKILLS
Python
CARRER OBJECTIVE
Entry-level data professional.
EXPERIENCE
Built a Python service.
CERTIFICATION
Certified training
175+ problems solved
"""

    sections = extract_sections(text)

    assert "B.Tech in CSE" in sections["education"]
    assert "candidate@example.com" in sections["contact"]
    assert "Entry-level" in sections["summary"]
    assert "175+ problems solved" in sections["achievements"]


def test_extract_sections_preserves_repeated_school_in_column_order() -> None:
    text = normalize_text(
        "Candidate Name\nname@example.com\n"
        "2022 - 2026\nCGPA - 9.38 (till 7\nSemester)\nth\n"
        "B.Tech in CSE\nExample University\n"
        "2022\nMarks (in percentage)-\n94.4%\nHigher Secondary (XII)\n"
        "Example High School\n"
        "2020\nMarks (in percentage)-\n93.4%\nSecondary (X)\n"
        "Example High School\nEDUCATION\nPROJECTS\nBuilt an analytics dashboard."
    )
    education = extract_sections(text)["education"]
    assert education.count("Example High School") == 2
    assert "2020\nMarks (in percentage)-\n93.4%\nSecondary (X)" in education


def test_extract_sections_returns_all_requested_ats_sections() -> None:
    sections = extract_sections("""Jane Doe | jane@example.com
CAREER OBJECTIVE
Data engineer
EDUCATION
B.Tech
EXPERIENCE
Internship
PROJECTS
Pipeline
SKILLS
Python
CERTIFICATIONS
AWS
ACHIEVEMENTS
Hackathon winner
""")

    assert set(sections) == {"contact", "summary", "education", "experience", "projects", "skills", "certifications", "achievements"}


def test_pdf_reader_retries_standard_extraction_when_layout_is_empty(tmp_path: Path, monkeypatch) -> None:
    pdf_path = tmp_path / "resume.pdf"
    pdf_path.write_bytes(b"fake")

    class Page:
        def extract_text(self, extraction_mode=None):
            return "" if extraction_mode == "layout" else "SUMMARY\nReadable resume text"

    fake_pypdf = SimpleNamespace(PdfReader=lambda _: SimpleNamespace(pages=[Page()]))
    monkeypatch.setitem(sys.modules, "pypdf", fake_pypdf)

    assert _read_pdf(pdf_path) == "SUMMARY\nReadable resume text"


def test_pdf_reader_prefers_logical_order_when_layout_interleaves_columns(tmp_path: Path, monkeypatch) -> None:
    pdf_path = tmp_path / "resume.pdf"
    pdf_path.write_bytes(b"fake")
    standard = "PROFILE SUMMARY\nUseful summary\nTECHNICAL SKILLS\nPython\nEDUCATION\nB.Tech"
    layout = "PROFILE SUMMARY                 TECHNICAL SKILLS\nUseful summary                    Python\nEDUCATION                         B.Tech"

    class Page:
        def extract_text(self, extraction_mode=None):
            return layout if extraction_mode == "layout" else standard

    fake_pypdf = SimpleNamespace(PdfReader=lambda _: SimpleNamespace(pages=[Page()]))
    monkeypatch.setitem(sys.modules, "pypdf", fake_pypdf)

    assert _read_pdf(pdf_path) == standard


def test_extract_sections_supports_academic_and_corporate_resume_headings() -> None:
    sections = extract_sections("""SAMEER SRIVASTAVA
PROFESSIONAL SUMMARY
Finance and analytics leader.
TEACHING VISION
Develop practitioners who bridge business and technology.
SUBJECTS AVAILABLE TO TEACH
Financial Modelling
ACADEMIC CREDENTIALS
MBA - Finance
KEY CORPORATE ACHIEVEMENTS
Improved efficiency by 37%.
CORPORATE EXPERIENCE (15+ YEARS)
Senior Manager | Example Ltd. | 2014 - 2025
CORE SKILLS & TOOLS
Domain Expertise
Financial Modelling • Market Sizing
Research and Consulting
Delivered client research.
""")

    assert "teaching_vision" in sections
    assert "teaching_subjects" in sections
    assert "MBA - Finance" in sections["education"]
    assert "37%" in sections["achievements"]
    assert "Senior Manager" in sections["experience"]
    assert "Financial Modelling" in sections["skills"]
    assert "publications" not in sections


def test_contact_recovery_scans_full_resume_without_treating_dates_as_phone_numbers() -> None:
    sections = extract_sections("""EDUCATION
2022 - 2026
B.Tech in Computer Science
Example College
LANGUAGES
English
Hindi
JANE DOE
+91 9083181985
Phone
jane@example.com
Email
CAREER OBJECTIVE
AI developer focused on applied machine learning.
""")

    assert "jane@example.com" in sections["contact"]
    assert "+91 9083181985" in sections["contact"]
    assert "2022 - 2026" not in sections["contact"]


def test_section_parser_does_not_treat_prose_beginning_with_experience_as_heading() -> None:
    sections = extract_sections("""PROFESSIONAL SUMMARY
Experience building accessible services for vulnerable communities.
CORE SKILLS
Case Management, Crisis Intervention, Community Outreach
""")

    assert "Experience building accessible services" in sections["summary"]
    assert "Case Management" in sections["skills"]


def test_domain_skills_and_roles_generalize_beyond_technology() -> None:
    text = """REGISTERED NURSE
CLINICAL SKILLS
Patient Care, Medication Administration, Wound Care, Discharge Planning
PROFESSIONAL MEMBERSHIPS
Indian Nursing Council
"""

    skills = extract_skills(text)
    assert {"Patient Care", "Medication Administration", "Wound Care", "Discharge Planning"} <= set(skills)
    assert "Registered Nurse" in extract_roles(text)
    assert "professional_memberships" in extract_sections(text)


def test_extract_sections_repairs_canva_column_order_by_content_type() -> None:
    sections = extract_sections("""ABOUT ME
Entry-level machine learning candidate.
LANGUAGES
JOYDIP PAUL
6289715644KolkataGitHubjoydippaul2004@gmail.com
EDUCATION
B.Tech | Computer Science | CGPA 8.27
English
Bengali
Hindi
PROJECTS
Churn Prediction
A Logistic Regression-based ML model for customer churn.
CORE SKILLS
LinkedIn
ACHIEVEMENTS
GATE Qualified (2025)
Programming: C, Java, Python.
AI/ML: Machine Learning, Deep Learning, TensorFlow.
Soft Skills: Problem-Solving, Collaboration
Image Deblurring using Deep Learning (Currently working)
A Deep Learning-based model for restoring images.
""")

    assert sections["languages"] == "English\nBengali\nHindi"
    assert "Programming: C, Java, Python." in sections["skills"]
    assert "AI/ML: Machine Learning" in sections["skills"]
    assert "Programming:" not in sections["achievements"]
    assert "Image Deblurring" in sections["projects"]
    assert "JOYDIP PAUL" in sections["contact"]


def test_contact_recovery_ignores_project_links_and_language_lines() -> None:
    sections = extract_sections("""PROJECTS
Classifier
Link - https://github.com/example/cat-dog-classifier
LANGUAGES
English
Hindi
JANE DOE
+91 9083181985
jane@example.com
KOLKATA
CAREER OBJECTIVE
Machine learning developer.
""")

    assert "https://github.com/example/cat-dog-classifier" not in sections["contact"]
    assert "English" not in sections["contact"]
    assert "KOLKATA" in sections["contact"]


def test_skill_category_labels_stay_inside_skills_and_internship_heading_is_experience() -> None:
    from job_hunting_agent.resume import extract_sections, extract_skills

    text = (
        "ANANYA SEN\nananya@example.com | +91 98765 43210\n"
        "Technical Skills\nLanguages: Python, C, C++, SQL\nTools: Git, Docker, Snowflake, Streamlit\n"
        "Internship Experience\nML Intern - Acme, Kolkata (Jan 2025 - Jun 2025)\n"
        "• Built a FastAPI service.\nLanguages\nEnglish, Hindi\n"
    )
    sections = extract_sections(text)
    assert "Languages: Python, C, C++, SQL" in sections["skills"]
    assert "ML Intern - Acme" in sections["experience"]
    assert "Python" not in sections.get("languages", "")
    skills = extract_skills(text)
    assert {"C", "Snowflake", "Streamlit"} <= set(skills)
    assert "Languages" not in skills


def test_spoken_language_label_inside_skills_goes_to_languages() -> None:
    from job_hunting_agent.resume import extract_sections

    sections = extract_sections("JOHN DOE\nSkills\nPython, SQL\nLanguages: English, Hindi\nTools: Git\nEducation\nB.Tech")
    assert sections["languages"] == "English, Hindi"
    assert "Tools: Git" in sections["skills"]


def test_wrapped_pdf_skill_line_does_not_invent_partial_skill() -> None:
    from job_hunting_agent.resume import extract_skills

    skills = extract_skills("JOHN DOE\nSkills\nPython, Kafka, AWS, Azure\nData Factory, Terraform\n")
    assert "Azure Data Factory" in skills
    assert "Data Factory" not in skills


def test_docx_list_bullets_are_preserved(tmp_path) -> None:
    from docx import Document
    from job_hunting_agent.resume import parse_resume

    document = Document()
    document.add_paragraph("JOHN DOE")
    document.add_paragraph("john@example.com | +91 98765 43210")
    document.add_heading("Projects", 1)
    document.add_paragraph("Helmet Detection")
    document.add_paragraph("Trained a YOLOv8 model on 4,000 images.", style="List Bullet")
    path = tmp_path / "resume.docx"
    document.save(path)
    resume = parse_resume(path)
    assert "• Trained a YOLOv8 model" in resume.sections["projects"]


def test_short_skills_are_not_inferred_from_prose() -> None:
    from job_hunting_agent.resume import extract_skills

    skills = extract_skills("Summary\nI go to work and grow net revenue with a C-level team.")
    assert not {"Go", "C", ".NET", "R"} & set(skills)
